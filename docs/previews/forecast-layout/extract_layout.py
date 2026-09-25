"""Read-only extraction of the supplied floor layout. No formula recalculation."""
import collections
import hashlib
import html
import json
from pathlib import Path

import openpyxl

SOURCE = Path(r"C:\Work Drive\APP\CNC OEE 참조파일\Setting CNC.xlsx")
DEST = Path(__file__).resolve().parent
sheet = openpyxl.load_workbook(SOURCE, data_only=True)["W39"]
machines = []
buildings = []
for building, first, last, offset in [("B", 4, 40, 0), ("A", 43, 87, 1600)]:
    count = 0
    for row in sheet.iter_rows(min_row=first, max_row=last, min_col=2, max_col=30):
        for cell in row:
            if not isinstance(cell.value, (int, float)) or int(cell.value) != cell.value or not 1 <= cell.value <= 800:
                continue
            label = sheet.cell(cell.row + 1, cell.column).value
            assert isinstance(label, str) and "-C" in label, (cell.coordinate, label)
            model, process = label.rsplit("-", 1)
            machines.append(dict(id=int(cell.value), building=building, cell=cell.coordinate,
                                 column=cell.column, row=cell.row, model=model, process=process,
                                 x=60 + (cell.column-2)*116,
                                 y=offset + 100 + (cell.row-first)*36,
                                 width=104, height=62))
            count += 1
    buildings.append(dict(id=building, count=count, x=28, y=offset+22, width=3400,
                          height=1490 if building == "B" else 1750))
assert len(machines) == 800 and {m['id'] for m in machines} == set(range(1, 801))
assert [b['count'] for b in buildings] == [448, 352]
data = dict(source=SOURCE.name, sheet="W39", sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            machines=sorted(machines, key=lambda m:m['id']), buildings=buildings,
            models=sorted({m['model'] for m in machines}))
(DEST/'layout-data.js').write_text('window.LAYOUT_DATA = '+json.dumps(data,ensure_ascii=False)+';\n',encoding='utf-8')
svg=['<svg xmlns="http://www.w3.org/2000/svg" width="3460" height="3420" viewBox="0 0 3460 3420">', '<rect width="3460" height="3420" fill="white"/>']
for b in buildings:
    svg.append(f'<text x="60" y="{b["y"]+40}" font-size="36" font-family="Arial">{b["id"]} / {b["count"]} machines</text>')
for m in machines:
    svg.append(f'<rect x="{m["x"]}" y="{m["y"]}" width="104" height="62" fill="#edf2fa" stroke="#8d9bb0"/>')
    svg.append(f'<text x="{m["x"]+6}" y="{m["y"]+26}" font-size="23" font-family="Arial">{m["id"]}</text>')
    svg.append(f'<text x="{m["x"]+6}" y="{m["y"]+48}" font-size="15" font-family="Arial">{html.escape(m["model"]+"-"+m["process"])}</text>')
svg.append('</svg>')
(DEST/'source-layout-reference.svg').write_text('\n'.join(svg),encoding='utf-8')
print(json.dumps({'result':'PASS','count':len(machines),'buildings':{b['id']:b['count'] for b in buildings},'sha256':data['sha256']}))
