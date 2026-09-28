"""Excel 배치 도면의 빨간 테두리(작업자 동선)에서 설비별 동선 번호·동선 쪽(L/R)을 뽑는다.

사용자 확정 2026-09-28: 빨간 선 = 작업자 동선이고, 그 양쪽 두 열이 마주본다. 동선은 가로 통로에서 끊긴다.
이 정보는 추천이 "한 동선 = 한 공정"을 지키는 데 쓰인다(supabase/migrations/20260928130000_layout_walkways.sql).

1공장(기본값, W39 시트 · B동 행 4–39 · A동 행 43–86):
  python scripts/extract-layout-walkways.py "../CNC OEE 참조파일/Setting CNC.xlsx" > walkways.json

2공장(Capa 시트 · B동 하나 · Excel 오타 보정, 사용자 확정 2026-09-28):
  python scripts/extract-layout-walkways.py "../CNC OEE 참조파일/ALV Setup.xlsx" --sheet Capa --count 350 \
    --building B:3:41:100 --fix F28=298 --fix F30=299 --fix F32=300 --fix F34=301 --fix F36=302

  --building CODE:TOP:BOTTOM[:Y0]  동 코드와 번호 행 범위. Y0 를 주면 도면 좌표(x, y, width, height)도 낸다 —
                                   1공장 좌표 규칙과 같다: x = 60 + 116·(열−2), y = Y0 + 36·(행−TOP), 104×62.
  --fix CELL=N                     번호 셀 값을 고쳐 읽는다(원본 Excel 은 그대로). 2공장 F28~F36 은 모두 302 로
                                   적혀 298~301 이 빠져 있다.
  --count N                        설비 대수(기본 800). 1..N 이 정확히 한 번씩 나와야 한다.

출력: {"<설비 번호>": {"walkway": "B-01-U", "side": "L", "cell": "B4", ...좌표}, ...}
  walkway = <동>-<동선 순번 2자리>-<U: 통로 위 | D: 통로 아래>, side = 동선의 왼쪽(L)/오른쪽(R) 열.
설비 번호 n 은 도면 좌표 마이그레이션과 같게 machines.name = 'CNC-' || lpad(n, 3, '0') 로 짝짓는다.
"""
import argparse
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter

DEFAULT_BUILDINGS = ('B:4:39', 'A:43:86')   # 1공장 W39 — 20260925120000 좌표 마이그레이션과 같은 구역


def is_red(side):
    return (side is not None and side.style and side.color is not None
            and isinstance(side.color.rgb, str) and side.color.rgb.upper().endswith('FF0000'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('path')
    ap.add_argument('--sheet', default='W39')
    ap.add_argument('--building', action='append')
    ap.add_argument('--fix', action='append', default=[])
    ap.add_argument('--count', type=int, default=800)
    args = ap.parse_args()

    buildings = []
    for spec in args.building or DEFAULT_BUILDINGS:
        parts = spec.split(':')
        buildings.append((parts[0], int(parts[1]), int(parts[2]), float(parts[3]) if len(parts) > 3 else None))
    fixes = {cell.upper(): int(n) for cell, n in (f.split('=') for f in args.fix)}

    ws_formula = openpyxl.load_workbook(args.path)[args.sheet]            # 테두리
    ws = openpyxl.load_workbook(args.path, data_only=True)[args.sheet]    # 값: 1공장 66–72, 74 는 수식 셀이라 계산값으로 읽는다

    def red_between(row, left_col):
        """left_col 과 left_col+1 사이 경계가 빨간가."""
        if left_col < 1:
            return False
        return is_red(ws_formula.cell(row=row, column=left_col).border.right) or \
            is_red(ws_formula.cell(row=row, column=left_col + 1).border.left)

    def number_at(row, col):
        coord = f'{get_column_letter(col)}{row}'
        if coord in fixes:
            return fixes[coord]
        v = ws.cell(row=row, column=col).value
        if isinstance(v, (int, float)) and float(v).is_integer() and 1 <= int(v) <= args.count:
            return int(v)
        return None

    out = {}
    for building, top, bottom, _ in buildings:
        # 설비 번호 셀 찾기(번호 셀 아래 행이 모델·공정 라벨)
        machines = []
        for row in range(top, bottom + 1):
            for col in range(1, ws.max_column + 1):
                n = number_at(row, col)
                if n is not None:
                    machines.append((n, row, col))
        for n, row, col in machines:
            if n in out:
                raise SystemExit(f'machine {n} appears twice ({out[n]["cell"]} and {get_column_letter(col)}{row}) — use --fix')
            # 동선 짝: 오른쪽 경계가 빨가면 (col, col+1), 왼쪽이 빨가면 (col-1, col).
            # 열 맨 위라 테두리가 없는 셀(예: 1공장 98)은 같은 열의 다른 설비가 가진 짝을 따른다.
            if red_between(row, col) or red_between(row + 1, col):
                left, side = col, 'L'
            elif red_between(row, col - 1) or red_between(row + 1, col - 1):
                left, side = col - 1, 'R'
            else:
                left, side = None, None
                for _, r2, c2 in machines:
                    if c2 != col or r2 == row:
                        continue
                    if red_between(r2, col):
                        left, side = col, 'L'
                        break
                    if red_between(r2, col - 1):
                        left, side = col - 1, 'R'
                        break
                if left is None:
                    raise SystemExit(f'machine {n} ({building} r{row}c{col}): no walkway found')
            # 가로 통로에서 동선이 끊긴다: 이 짝 경계의 빨간 행을 연속 구간으로 나눠 몇 번째 구간인지 본다.
            red_rows = [r for r in range(top, bottom + 1) if red_between(r, left)]
            runs, start, prev = [], None, None
            for r in red_rows:
                if start is None:
                    start = prev = r
                elif r == prev + 1:
                    prev = r
                else:
                    runs.append((start, prev))
                    start = prev = r
            if start is not None:
                runs.append((start, prev))
            section = min(range(len(runs)), key=lambda i: 0 if runs[i][0] <= row + 1 and row <= runs[i][1]
                          else min(abs(row - runs[i][0]), abs(row - runs[i][1])))
            out[n] = {'building': building, 'pair': left, 'run_top': runs[section][0], 'side': side,
                      'cell': f'{get_column_letter(col)}{row}', 'row': row, 'col': col}

    # 동별 가로 통로: 어떤 짝에서도 빨간 테두리가 없는 행 중 동 구역 안쪽에 있는 행
    result = {}
    for building, top, bottom, y0 in buildings:
        items = {n: v for n, v in out.items() if v['building'] == building}
        pairs = sorted({v['pair'] for v in items.values()})
        aisle_rows = [r for r in range(top + 2, bottom - 1)
                      if not any(red_between(r, p) for p in pairs)
                      and any(red_between(r - 1, p) for p in pairs) and any(red_between(r + 1, p) for p in pairs)]
        if len(aisle_rows) != 1:
            raise SystemExit(f'{building}: expected exactly one aisle row, got {aisle_rows}')
        aisle = aisle_rows[0]
        for n, v in items.items():
            index = pairs.index(v['pair']) + 1
            entry = {'walkway': f"{building}-{index:02d}-{'U' if v['run_top'] < aisle else 'D'}",
                     'side': v['side'], 'cell': v['cell']}
            if y0 is not None:
                entry.update({'building': building, 'x': 60 + 116 * (v['col'] - 2), 'y': y0 + 36 * (v['row'] - top),
                              'width': 104, 'height': 62})
            result[n] = entry

    if len(result) != args.count:
        missing = [n for n in range(1, args.count + 1) if n not in result]
        raise SystemExit(f'expected {args.count} machines, got {len(result)}; missing {missing}')
    json.dump({str(n): result[n] for n in sorted(result)}, sys.stdout, ensure_ascii=False, indent=0)


if __name__ == '__main__':
    main()
